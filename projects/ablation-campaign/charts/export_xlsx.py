#!/usr/bin/env python3
"""Export the master table as an Excel workbook with one tab per evaluation set.

    uv run --no-project --with openpyxl python export_xlsx.py [--out results_master.xlsx] [--master PATH]

The master is only read. Tabs: ``README`` (what each tab and metric is), then one tab per set in the registry, test
sets first. In a set's tab each row is one model under one prompt format and each column one ``metric`` x ``unit``
(a language or an aggregate); cells the master leaves empty stay empty. Values are written as numbers, not text.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import master


SPLIT_ORDER = {"test": 0, "dev": 1, "reported": 2}
TAB_MAX = 31  # Excel's limit on a sheet name
HEADER_FILL = "DDE6F0"


def _tab_name(name: str, used: set[str]) -> str:
    tab = name[:TAB_MAX]
    if tab in used:
        raise master.MasterError(f"two sets share the tab name {tab!r}")
    used.add(tab)
    return tab


def _number(text: str) -> float | str:
    try:
        return float(text)
    except ValueError:
        return text


def build_workbook(paths: master.Paths):
    """Return an openpyxl ``Workbook`` for the master at *paths*."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    table = master.read_table(paths.master)
    sets = master.read_sets(paths.sets)
    by_set: dict[str, list[str]] = {}
    for set_name, fmt in sets:
        by_set.setdefault(set_name, []).append(fmt)
    order = sorted(by_set, key=lambda s: (SPLIT_ORDER.get(sets[(s, by_set[s][0])]["split"], 9), s))

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor=HEADER_FILL)
    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    used = {"README"}
    readme.append(["tab", "task", "split", "formats", "metrics", "description (first format)"])
    for cell in readme[1]:
        cell.font = bold
        cell.fill = head_fill

    meta_cols = ["label", "model", "fmt", "kind", "recipe"]
    for set_name in order:
        fmts = sorted(by_set[set_name])
        columns = [c for c in table.data_columns if master.parse_column(c)[0] == set_name]
        columns.sort(key=master.column_sort_key)
        tab = _tab_name(set_name, used)
        info = sets[(set_name, fmts[0])]
        readme.append([tab, info["task"], info["split"], " ".join(fmts), info["metrics"], info["description"]])
        ws = wb.create_sheet(tab)

        # One column per (metric, unit) shared by every fmt, so rows of different formats line up.
        keys = sorted(
            {(master.parse_column(c)[2], master.parse_column(c)[3]) for c in columns},
            key=lambda mu: (master.METRIC_ORDER.get(mu[0], 99), mu[0], master.UNIT_RANK.get(mu[1], 99), mu[1]),
        )
        ws.append(meta_cols + [metric for metric, _ in keys])
        ws.append([""] * len(meta_cols) + [unit for _, unit in keys])
        for row in ws.iter_rows(min_row=1, max_row=2):
            for cell in row:
                cell.font = bold
                cell.fill = head_fill
                cell.alignment = Alignment(horizontal="center")
        for key, row in sorted(table.rows.items(), key=lambda kv: (int(kv[1].get("order") or 999), kv[0])):
            for fmt in fmts:
                values = [row.get(f"{set_name}.{fmt}.{metric}.{unit}", "") for metric, unit in keys]
                if not any(values):
                    continue
                ws.append([row["label"], key, fmt, row["kind"], row["recipe"], *map(_number, values)])
        ws.freeze_panes = ws.cell(row=3, column=len(meta_cols) + 1)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 24
        for i in range(len(meta_cols) + 1, len(meta_cols) + len(keys) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 9
        digits = "0.000" if set_name.startswith("selection") else "0.00"
        for row in ws.iter_rows(min_row=3, min_col=len(meta_cols) + 1):
            for cell in row:
                cell.number_format = digits
    readme.column_dimensions["A"].width = 28
    readme.column_dimensions["F"].width = 120
    return wb


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--master", help="path of results_master.csv (default: the one next to this script)")
    parser.add_argument("--out", default="results_master.xlsx", help="workbook to write (default: %(default)s)")
    args = parser.parse_args(argv)
    try:
        wb = build_workbook(master.resolve_paths(args.master))
    except ImportError:
        print(
            "openpyxl is missing; run with `uv run --no-project --with openpyxl python export_xlsx.py`",
            file=sys.stderr,
        )
        return 2
    except master.MasterError as err:
        print(err, file=sys.stderr)
        return 1
    out = Path(args.out)
    wb.save(out)
    print(f"{out}: {len(wb.sheetnames) - 1} set tabs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
